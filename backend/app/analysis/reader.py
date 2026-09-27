import csv
import math
import zipfile
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.errors import AppError

MAX_CELL_LENGTH = 10_000


def normalize_headers(headers: list) -> tuple[list[str], list[dict]]:
    used: set[str] = set()
    names, mapping = [], []
    for index, raw in enumerate(headers):
        original = "" if raw is None else str(raw)
        base = original.strip() or f"column_{index + 1}"
        if len(base) > 200:
            raise AppError("HEADER_TOO_LONG", "Ustun nomi 200 belgidan oshmasligi kerak.")
        name, suffix = base, 2
        while name in used:
            name = f"{base}_{suffix}"
            suffix += 1
        used.add(name)
        names.append(name)
        mapping.append({"original": original, "name": name})
    return names, mapping


def read_frame(path: Path, options: dict, limits: dict) -> tuple[pd.DataFrame, dict]:
    warnings = []
    sheets: list[str] = []
    selected_sheet = None
    rows = []

    def append_row(row):
        if len(rows) >= limits["max_rows"]:
            raise AppError("ROW_LIMIT", f"Jadval {limits['max_rows']:,} qatordan oshmasligi kerak.")
        if len(row) != len(headers):
            raise AppError("INCONSISTENT_COLUMNS", "Qatorlardagi ustun soni header bilan mos emas.")
        if any(isinstance(value, str) and len(value) > MAX_CELL_LENGTH for value in row):
            raise AppError(
                "CELL_TOO_LONG", "Bitta katakdagi matn 10 000 belgidan oshmasligi kerak."
            )
        rows.append([None if value == "" else value for value in row])

    def check_header():
        if not headers:
            raise AppError("EMPTY_FILE", "Faylda jadval topilmadi.")
        if len(headers) > limits["max_columns"]:
            raise AppError("COLUMN_LIMIT", "Jadvalda ustunlar soni limitdan oshgan.")

    if path.suffix == ".csv":
        try:
            with path.open(encoding="utf-8-sig", newline="") as stream:
                sample = stream.read(8192)
                if "\x00" in sample:
                    raise AppError("INVALID_CSV", "CSV matn fayli bo‘lishi kerak.")
                stream.seek(0)
                delimiter = options.get("delimiter")
                if delimiter is None:
                    try:
                        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
                    except csv.Error:
                        delimiter = ","
                        warnings.append(
                            "Ajratgich aniqlanmadi; vergul ishlatildi. Previewni tekshiring."
                        )
                reader = csv.reader(stream, delimiter=delimiter, strict=True)
                headers = next(reader, [])
                check_header()
                for row in reader:
                    if not row:
                        continue
                    append_row(row)
        except UnicodeDecodeError as exc:
            raise AppError("ENCODING", "CSV faylni UTF-8 formatida saqlab qayta yuklang.") from exc
        except csv.Error as exc:
            raise AppError(
                "INVALID_CSV", "CSV tuzilishi noto‘g‘ri yoki katak hajmi juda katta."
            ) from exc
    else:
        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(item.file_size for item in entries) > 200 * 1024**2:
                    raise AppError("XLSX_LIMIT", "Excel faylining ochilgan hajmi limitdan oshgan.")
                if any("vbaproject" in item.filename.lower() for item in entries):
                    raise AppError("MACROS_NOT_ALLOWED", "Makrosli Excel fayli qabul qilinmaydi.")
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                sheets = workbook.sheetnames
                if len(sheets) > 20:
                    raise AppError("SHEET_LIMIT", "Ko‘pi bilan 20 sheetga ruxsat beriladi.")
                selected_sheet = options.get("sheet") or sheets[0]
                if selected_sheet not in sheets:
                    raise AppError("INVALID_SHEET", "Tanlangan sheet faylda mavjud emas.")
                sheet = workbook[selected_sheet]
                if sheet.max_column and sheet.max_column > limits["max_columns"]:
                    raise AppError(
                        "COLUMN_LIMIT", "Excel sheetidagi ustunlar soni limitdan oshgan."
                    )
                # Do not trust a workbook's cached row dimensions (some producers lie).
                sheet.reset_dimensions()
                iterator = sheet.iter_rows(values_only=True)
                headers = list(next(iterator, ()))
                check_header()
                for values in iterator:
                    row = list(values)
                    if len(row) > len(headers) and any(v is not None for v in row[len(headers) :]):
                        raise AppError(
                            "INCONSISTENT_COLUMNS", "Headerdan tashqarida qiymatlar bor."
                        )
                    row = row[: len(headers)] + [None] * max(0, len(headers) - len(row))
                    append_row(row)
            finally:
                workbook.close()
            warnings.append("Excel formulalari bajarilmaydi; faqat saqlangan qiymatlar o‘qiladi.")
            if len(sheets) > 1 and not options.get("sheet"):
                warnings.append(
                    f"Birinchi sheet tanlandi: {selected_sheet}. Boshqa sheetni tanlashingiz mumkin."
                )
        except AppError:
            raise
        except Exception as exc:
            raise AppError(
                "INVALID_XLSX", "Excel faylini o‘qib bo‘lmadi. Oddiy .xlsx fayl yuklang."
            ) from exc

    if not rows:
        raise AppError("EMPTY_DATASET", "Jadvalda headerdan tashqari ma’lumot yo‘q.")
    names, mapping = normalize_headers(headers)
    frame = pd.DataFrame(rows, columns=names)
    # Convert only wholly numeric columns, preserving identifiers such as 00123.
    for name in names:
        series = frame[name]
        values = series.dropna()
        if values.empty or pd.api.types.is_datetime64_any_dtype(series):
            continue
        if values.map(lambda value: isinstance(value, (date, datetime, bool))).any():
            continue
        text_values = values.astype(str)
        leading_zero = text_values.str.match(r"^[+-]?0\d+").any()
        converted = pd.to_numeric(series, errors="coerce")
        if not leading_zero and converted.notna().sum() == len(values):
            # Integers beyond IEEE-754 exact range are usually identifiers.
            if (converted.dropna().abs() > 2**53 - 1).any():
                frame[name] = series.map(lambda value: None if pd.isna(value) else str(value))
                warnings.append(
                    f"{name}: katta sonlar aniqlikni saqlash uchun matn sifatida olindi."
                )
            else:
                frame[name] = converted
    if any(item["original"] != item["name"] for item in mapping):
        warnings.append("Bo‘sh yoki takroriy ustun nomlari normallashtirildi; mapping saqlandi.")
    return frame, {
        "sheets": sheets,
        "selected_sheet": selected_sheet,
        "delimiter": delimiter if path.suffix == ".csv" else None,
        "column_mapping": mapping,
        "warnings": warnings,
    }


def json_value(value):
    if value is None or pd.isna(value):
        return None
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def table(frame: pd.DataFrame, limit: int = 100) -> dict:
    return {
        "columns": list(frame.columns),
        "rows": [
            [json_value(value) for value in row]
            for row in frame.head(limit).itertuples(index=False, name=None)
        ],
        "total_rows": len(frame),
        "truncated": len(frame) > limit,
    }
